import io
import docx
from fastapi.testclient import TestClient
from app import app

client = TestClient(app)

def test_knowledge_upload():
    print("Testing /api/knowledge/upload with DOCX file...")
    doc = docx.Document()
    doc.add_heading("GHL Custom Knowledge Test", 0)
    doc.add_paragraph("This document explains how to set up headless portals with GoHighLevel OAuth 2.0 and webhooks.")
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_bytes = doc_io.getvalue()

    files = {
        "file": ("ghl_guide.docx", doc_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    }

    response = client.post("/api/knowledge/upload", files=files)
    print("Status code:", response.status_code)
    data = response.json()
    print("Response JSON:", data)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}: {data}"
    assert data.get("status") == "success"
    assert data.get("filename") == "ghl_guide.docx"
    assert data.get("chunks_indexed", 0) >= 1
    print("✅ /api/knowledge/upload verified successfully!")

if __name__ == "__main__":
    test_knowledge_upload()
